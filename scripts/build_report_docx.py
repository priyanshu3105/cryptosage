#!/usr/bin/env python3
"""One-off generator for CryptoSage_Report.docx (project root)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
def add_mono_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "CryptoSage_Report.docx"

    doc = Document()

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Technical Report: CryptoSage")
    r.bold = True
    r.font.size = Pt(22)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        "Author: Priyanshu Chatterjee",
        "Date: 29 April 2026",
        "Course / Module: (please fill in)",
    ):
        meta.add_run(line + "\n")

    doc.add_page_break()

    # Abstract
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "CryptoSage is a full-stack web application for cryptocurrency and DeFi literacy, "
        "portfolio tracking, and journal-based trade logging. It integrates live market data "
        "(CoinGecko), DeFi protocol metrics (DefiLlama), and an educational AI assistant (Groq) "
        "with explicit safeguards against financial advice. This report describes objectives, "
        "development methodology, system architecture, implementation details, use of "
        "AI-assisted coding, deployment status, and conclusions with future work."
    )

    # 1 Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "CryptoSage addresses fragmentation across price trackers, journals, DeFi dashboards, "
        "and generic chat tools by unifying portfolio tracking, market exploration, DeFi browsing, "
        "and a domain-restricted assistant in one application."
    )
    doc.add_heading("Objectives", level=2)
    for item in (
        "Deliver secure authentication and protected application routes.",
        "Persist users, portfolios, holdings, journal logs, and chat sessions in MongoDB.",
        "Expose a consistent REST API with validation and interactive API documentation (Swagger).",
        "Integrate external market and DeFi APIs with caching, retries, and graceful degradation.",
        "Provide an AI assistant focused on education and risk awareness, not investment advice.",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Context and motivation", level=2)
    doc.add_paragraph(
        "Retail crypto users need trustworthy structure around data and explanations. "
        "CryptoSage emphasizes literacy, journaling, and safe assistant behaviour rather than "
        "speculative guidance."
    )

    # 2 Development Strategy
    doc.add_heading("2. Development Strategy", level=1)
    doc.add_heading("Methodology", level=2)
    doc.add_paragraph(
        "An iterative, feature-first workflow was used: scaffolding and routing; authentication; "
        "portfolio and journal persistence; market and DeFi integrations; AI chat with safety "
        "and fallback logic; then UX polish and error handling."
    )
    doc.add_heading("Planning decisions", level=2)
    for item in (
        "Domain-separated backend modules (auth, portfolio, market, defi, chat).",
        "Contract-first APIs with Zod validation on request bodies, queries, and params.",
        "Service and repository layering to isolate business logic from persistence.",
        "Security and rate limiting treated as first-class requirements, not afterthoughts.",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Justification", level=2)
    doc.add_paragraph(
        "This structure reduces coupling between the React frontend and Express backend, "
        "speeds parallel development, and makes AI-adjacent and auth-sensitive code easier to review."
    )

    # 3 Implementation
    doc.add_heading("3. Implementation", level=1)
    doc.add_heading("3.1 Technology stack", level=2)
    doc.add_paragraph(
        "Frontend: React 18, TypeScript, Vite, React Router, TanStack Query, Tailwind/shadcn UI, "
        "Recharts, Framer Motion. Backend: Node.js, Express 5, TypeScript, Mongoose, Zod, JWT, "
        "bcrypt, Axios, Pino logging. Database: MongoDB (Atlas-compatible). External services: "
        "CoinGecko, DefiLlama, Groq chat completions."
    )

    doc.add_heading("3.2 System architecture", level=2)
    doc.add_paragraph(
        "Figure 1 (high-level architecture). Export from mermaid.live as PNG/SVG and insert here if required."
    )
    add_mono_block(
        doc,
        "flowchart LR\n"
        "  U[User Browser] --> FE[React Frontend]\n"
        "  FE -->|REST /api| BE[Express Backend]\n"
        "  BE --> DB[(MongoDB)]\n"
        "  BE --> CG[CoinGecko]\n"
        "  BE --> DL[DefiLlama]\n"
        "  BE --> GQ[Groq]",
    )

    doc.add_heading("3.3 Backend structure", level=2)
    doc.add_paragraph(
        "Express application factory wires Helmet, CORS, JSON body limits, request IDs "
        "(x-request-id), global and route-specific rate limits, Swagger UI at /api-docs, "
        "and route prefixes: /api/auth, /api/portfolio, /api/market, /api/defi, /api/chat, "
        "plus /api/health. Centralised not-found and error handlers return structured errors "
        "with requestId."
    )

    doc.add_heading("3.4 Data model", level=2)
    doc.add_paragraph("Figure 2 (entity relationships). Mermaid source:")
    add_mono_block(
        doc,
        "erDiagram\n"
        "  USER ||--|| PORTFOLIO : owns\n"
        "  USER ||--o{ PORTFOLIO_LOG : creates\n"
        "  USER ||--o{ CHAT_SESSION : has\n"
        "  PORTFOLIO ||--o{ HOLDING : contains",
    )

    doc.add_heading("3.5 Key features", level=2)
    for item in (
        "Authentication: register, login, JWT, bcrypt password hashing, /auth/me and account lifecycle endpoints.",
        "Portfolio: per-user portfolio with holdings; journal buy/sell replay can rebuild holdings with average cost basis.",
        "Market: top coins, price and history endpoints backed by CoinGecko with in-memory TTL cache and retries.",
        "DeFi: protocol list and detail/TVL history from DefiLlama with defensive parsing for varying payload shapes.",
        "Chat: topic gate, lightweight keyword-based retrieval context, structured JSON from the model with parse/retry, policy guard, session append in MongoDB when sessionId is provided.",
        "Portfolio insights: authenticated POST /chat/portfolio-insights combines logs, optional holdings, optional market snapshot.",
        "Frontend: protected routes, portfolio dashboard, market and DeFi explorers with charts, journal with AI side panel, local calculator tools, assistant widget with offline fallbacks.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3.6 AI assistant pipeline", level=2)
    doc.add_paragraph("Figure 3 (sequence). Mermaid source:")
    add_mono_block(
        doc,
        "sequenceDiagram\n"
        "  participant UI as Frontend\n"
        "  participant API as Backend\n"
        "  participant TG as Topic gate\n"
        "  participant RAG as Retrieval\n"
        "  participant LLM as Groq\n"
        "  participant PG as Policy guard\n"
        "  UI->>API: POST /chat\n"
        "  API->>TG: classify\n"
        "  alt blocked\n"
        "    API-->>UI: refusal\n"
        "  else allowed\n"
        "    API->>RAG: context\n"
        "    API->>LLM: structured prompt\n"
        "    LLM-->>API: JSON\n"
        "    API->>PG: validate\n"
        "    API-->>UI: answer or refusal\n"
        "  end",
    )

    doc.add_heading("3.7 Design decisions and trade-offs", level=2)
    doc.add_paragraph(
        "Strengths: modular routes, validated inputs, consistent API envelopes, resilience patterns "
        "for third-party APIs, and deliberate AI safety boundaries. Trade-offs: in-memory cache "
        "and custom rate limit stores do not scale across multiple server instances without a "
        "shared store (e.g. Redis). Automated tests are minimal in the current tree; production "
        "CI/CD is not fully wired at repository root."
    )

    # 4 Vibe coding
    doc.add_heading("4. Vibe Coding Strategy", level=1)
    doc.add_paragraph(
        "AI-assisted tools (e.g. Cursor, Copilot-class assistants) were used to accelerate UI "
        "scaffolding, TypeScript boilerplate, and exploration of edge cases. Human review remained "
        "essential for authentication, data integrity, prompt safety, and API contracts. "
        "Iteration followed: prototype, validate with real API behaviour, then harden validation, "
        "errors, and fallbacks."
    )

    # 5 Deployment
    doc.add_heading("5. Deployment", level=1)
    doc.add_paragraph(
        "Local development: frontend via Vite (default port 8080) with /api proxied to the backend "
        "port read from backend/.env; backend via npm run dev (PORT from env, default 4000). "
        "MongoDB connection string is required in .env (MONGODB_URI); backend README notes Atlas "
        "and IP allowlisting. Swagger is served at /api-docs for reproducible API exploration."
    )
    doc.add_paragraph(
        "A full production CI/CD pipeline is not yet established in the repository; recommended "
        "future work includes lint, test, and build stages plus staged deployment."
    )

    # 6 Conclusions
    doc.add_heading("6. Conclusions and Future Work", level=1)
    doc.add_paragraph(
        "CryptoSage meets its core goals as an integrated literacy and tracking platform with "
        "defensible AI integration. The strongest aspects are clear layering, input validation, "
        "and breadth of user-facing features. The main gaps are deeper automated testing, "
        "distributed caching and rate limiting, and formal release automation."
    )
    doc.add_heading("Future work", level=2)
    for item in (
        "Expand unit, integration, and end-to-end tests for auth, portfolio, and chat paths.",
        "Introduce Redis (or equivalent) for shared rate limits and cache in multi-instance deployments.",
        "Add CI/CD (GitHub Actions or GitLab CI) with lint, test, and build artifacts.",
        "Richer portfolio analytics and assistant evaluation telemetry.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    # References
    doc.add_heading("References", level=1)
    for item in (
        "CoinGecko API — https://www.coingecko.com/en/api",
        "DefiLlama API — https://defillama.com/docs/api",
        "Groq API — https://console.groq.com/docs",
        "Repository: cryptosage (frontend and backend package.json scripts: npm run dev, npm run build).",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
